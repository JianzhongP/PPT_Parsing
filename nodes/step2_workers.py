"""Step 2: 各类型元素的并行分析Agent"""

import os

from state import PageElement, ElementInsight, PPTPageState
from openai import AzureOpenAI
from pptx import Presentation
from nodes.step1_global_analysis import ImageUtils
from nodes.type_normalization import normalize_element_type
from typing import Dict, Optional, Any
from markitdown import MarkItDown
from vlm_client import VLMClient

try:
    # 复用复杂 pipeline 中已经实现的合合(TextIn)客户端
    from nodes.complex_pipeline.phase3_expert_extraction import TextInAPIClient
except Exception:
    TextInAPIClient = None  # type: ignore

md_parser = MarkItDown()


_TEXTIN_CLIENT = None


def _bbox_iou(a: list[int], b: list[int]) -> float:
    ay1, ax1, ay2, ax2 = a
    by1, bx1, by2, bx2 = b
    inter_x1 = max(ax1, bx1)
    inter_y1 = max(ay1, by1)
    inter_x2 = min(ax2, bx2)
    inter_y2 = min(ay2, by2)
    iw = max(0, inter_x2 - inter_x1)
    ih = max(0, inter_y2 - inter_y1)
    inter = iw * ih
    if inter <= 0:
        return 0.0
    area_a = max(0, ax2 - ax1) * max(0, ay2 - ay1)
    area_b = max(0, bx2 - bx1) * max(0, by2 - by1)
    union = area_a + area_b - inter
    return float(inter / union) if union > 0 else 0.0


def _pick_native_table_for_element(element: PageElement, native_tables: list[dict]) -> Optional[dict]:
    if not native_tables:
        return None

    el_box = getattr(getattr(element, "bbox", None), "box_2d", None)
    if not el_box or len(el_box) != 4:
        return native_tables[0]

    best_table = None
    best_iou = 0.0
    for table in native_tables:
        tb = table.get("bbox")
        if isinstance(tb, list) and len(tb) == 4:
            iou = _bbox_iou([int(v) for v in el_box], [int(v) for v in tb])
            if iou > best_iou:
                best_iou = iou
                best_table = table

    if best_table is not None and best_iou >= 0.1:
        return best_table
    return native_tables[0]


def _get_textin_client() -> Optional[Any]:
    """Lazy init 合合(TextIn)客户端。

    通过环境变量读取：
    - TEXTIN_APP_ID / TEXTIN_SECRET_CODE
    - 或 TEXTIN_API_KEY="app_id:secret_code"
    """
    global _TEXTIN_CLIENT
    if _TEXTIN_CLIENT is not None:
        return _TEXTIN_CLIENT

    if TextInAPIClient is None:
        _TEXTIN_CLIENT = None
        return None

    app_id = os.getenv("TEXTIN_APP_ID", "").strip()
    secret_code = os.getenv("TEXTIN_SECRET_CODE", "").strip()
    api_key = os.getenv("TEXTIN_API_KEY", "").strip()

    try:
        if app_id and secret_code:
            _TEXTIN_CLIENT = TextInAPIClient(app_id=app_id, secret_code=secret_code)
        elif api_key:
            _TEXTIN_CLIENT = TextInAPIClient(api_key=api_key)
        else:
            _TEXTIN_CLIENT = TextInAPIClient()
    except Exception:
        _TEXTIN_CLIENT = None
    return _TEXTIN_CLIENT


# ============================================================================
# 1. 纯文本 Worker
# ============================================================================

def convert_slide_to_markdown(ppt_path: str, page_index: int) -> str:
    """
    直接读取 PPTX 底层 XML 将特定页面转换为 Markdown
    模拟 MarkItDown 的核心逻辑，但支持单页操作
    """
    try:
        prs = Presentation(ppt_path)
        
        # 越界保护
        if page_index >= len(prs.slides):
            return ""
            
        slide = prs.slides[page_index]
        md_lines = []
        
        # 1. 提取标题
        if slide.shapes.title and slide.shapes.title.text:
            md_lines.append(f"# {slide.shapes.title.text.strip()}")
            md_lines.append("")
            
        # 2. 遍历形状提取文本 (按阅读顺序大致排序)
        # 简单的排序策略：先上后下，先左后右
        shapes = sorted(
            [s for s in slide.shapes if hasattr(s, "text") and s.text.strip()], 
            key=lambda x: (x.top, x.left)
        )
        
        for shape in shapes:
            # 跳过标题（已处理）
            if shape == slide.shapes.title:
                continue
                
            text = shape.text.strip()
            if not text:
                continue
            
            # 简单的层级处理（如果原本是项目符号）
            # 实际的 XML 解析可以更复杂，这里做简化处理
            md_lines.append(text)
            md_lines.append("")
            
        # 3. 提取备注 (Speaker Notes) - 这在图片模式下是绝对拿不到的
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                md_lines.append("---")
                md_lines.append(f"> **演讲者备注**: {notes}")
                
        return "\n".join(md_lines)
        
    except Exception as e:
        print(f"❌ [TextWorker] PPTX XML 解析失败: {e}")
        return ""

def node_text_worker(state: PPTPageState, llm_client: AzureOpenAI, debug_logger=None) -> dict:
    """处理纯文本页面的节点"""
    page_index = state.page_index
    
    print(f"[TextWorker] 基于 XML 解析页面 {page_index}")
    
    # 记录步骤开始
    if debug_logger:
        debug_logger.log_step_start(
            page_index,
            "Step 2: Text Worker",
            input_data={
                "page_index": page_index,
                "ppt_path": state.ppt_path,
                "global_analysis": state.global_analysis.section_title if state.global_analysis else None
            },
            previous_step="Step 2: Router"
        )
    
    # 1. XML 原生解析
    xml_content = convert_slide_to_markdown(state.ppt_path, state.page_index)
    
    # 2. VLM 视觉提取
    vlm_content = state.global_analysis.extracted_text
    
    # 构造 Insight
    insight = ElementInsight(
        element_id="full_text_page",
        element_type="text_page",
        key_insight=f"XML源文本内容", 
        data_evidence=xml_content,
        confidence=1.0,
        status="pass",
        hypothesis_verification=f"VLM辅助文本: {vlm_content[:100]}..."
    )
    
    # 记录步骤结束
    if debug_logger:
        debug_logger.log_step_end(
            page_index,
            "Step 2: Text Worker",
            output_data={
                "element_insights_count": 1,
                "element_id": insight.element_id,
                "element_type": insight.element_type,
                "confidence": insight.confidence,
                "key_insight_preview": insight.key_insight[:100]
            },
            status="success"
        )
    
    return {"element_insights": [insight]}



class ElementAnalysisAgent:
    """元素分析Agent基类
    
    职责：
    1. 接收第一步的输出（全局分析 + 元素信息）
    2. 对指定元素进行深度分析
    3. 提取与全局上下文相关的核心信息
    4. 不输出完整的元素信息，只输出汇报者想表达的核心内容
    """
    
    def __init__(self, vlm_client: VLMClient, element_type: str):
        self.vlm_client = vlm_client
        self.element_type = element_type
        from config import get_config
        self.model = get_config().vlm_runtime_model_name
    
    def get_system_prompt(self) -> str:
        """获取该类型元素的专用prompt"""
        prompts = {
            "km_curve": """你是一个肿瘤临床数据专家。
任务：分析 Kaplan-Meier 生存曲线。
重点提取：
1. **组别信息**：实验组和对照组分别是什么？
2. **核心统计量**：Hazard Ratio (HR) 值及其置信区间 (CI)，P 值 (P-value)。
3. **中位生存期**：mOS 或 mPFS 具体数值。
4. **结论**：曲线是否分离？哪一组表现更优？
输出：结构化的统计数据和临床结论。""",

            "forest_plot": """你是一个生物统计专家。
任务：分析森林图 (Forest Plot)。
重点提取：
1. **分析目的**：这是亚组分析还是多研究汇总？
2. **显著性**：哪些亚组/研究显示出了统计学显著性（置信区间未跨过 1）？
3. **趋势**：整体效应倾向于哪一方（获益 vs 风险）？
输出：关键的显著亚组和整体效应值。""",

            "molecular_structure": """你是一个药物化学家。
任务：识别化学分子结构。
重点提取：
1. **母核结构**：是什么类型的骨架（如单抗、ADC linker、小分子母核）？
2. **关键基团**：有什么特殊的取代基或偶联位点？
3. **药物名称**：如果图中包含名称（如代码或通用名），请提取。
输出：分子结构特征描述。""",

            "flowchart": """你是一个临床试验设计或生物学专家。
任务：解析流程图或机制通路图，必须**穷尽式地提取**图中所有文字信息。
重点提取：
1. **全局总结**：该流程图/机制图的核心目的。
2. **全量节点与逻辑**：按照从上到下、从左到右或逻辑先后顺序（根据箭头指向判断），提取每个阶段/节点（包含筛选、入组、各分支条件、处理方式、具体数据等）中的**所有详细文字内容**。绝不能省略、概括或遗漏任何细节文本。
输出：核心结论 + 清晰、详尽的流程步骤描述（使用Markdown结构化输出）。""",

            "data_chart": """你是一个商业数据分析师。
任务：分析数据图表（柱状、折线、饼图）。
重点提取：
1. **趋势**：增长率、下降幅度。
2. **极值**：最大值、最小值、关键拐点。
3. **对比**：不同年份或竞品之间的倍数关系。
输出：核心数据趋势和商业洞察。""",
            
            "table": """你是一个精细的审计员。
任务：从表格中提取**最关键的数值**和对比。
重点：
- 哪一行/列最重要？
- 有什么显著的数值差异或统计指标（如p值、增长率）？
- 这个信息如何支持全局结论？
输出：只列出关键指标和结论，不要全部罗列。""",

            "formula": """你是一个数学公式解析专家。
任务：识别并解释图片中的核心公式。
重点提取：
1. 公式本体（可用 LaTeX 或原文）。
2. 各符号/变量的含义（若可判断）。
3. 该公式在当前页面语境中的作用。
输出：简洁、结构化的公式解读。""",
            
            "image": """你是一个视觉内容理解专家。
任务：理解图片中的**核心内容**，并**穷尽式地提取**图片中的所有有效文本。
重点：
1. **核心总结**：这张图的主题和核心信息。
2. **全量文本结构化**：有条理地识别并输出图片中的所有文字、标注、说明、步骤和数据，按照空间位置或逻辑关系进行组织。不要省略和自己捏造！
输出：核心描述 + 详尽的图内全量文本结构化呈现。""",
            
            "unknown": """请仔细分析这个元素，**全面且有条理地提取图片中出现的所有文字、数据和模块信息**。在总结核心结论的同时，提供结构化的全量文本描述，不要遗漏关键信息。"""
        }
        # 模糊匹配策略：防止 VLM 输出 "km_curve_v2" 导致匹配失败
        for key in prompts:
            if key in self.element_type.lower():
                return prompts[key]
        
        return prompts["unknown"]
    
    def analyze_element(self, 
                       element: PageElement, 
                       global_summary: str, 
                       original_image_path: str,
                       hypothesis: str = "",
                       native_page_evidence: Optional[dict] = None) -> ElementInsight:
        """分析单个元素"""
        print(f"[元素Agent] 分析 {self.element_type} 元素: {element.element_id}")
        
        # 尝试裁剪出该元素的图片
        target_image_path = original_image_path
        if element.bbox and element.bbox.box_2d:
            target_image_path = ImageUtils.crop_image(original_image_path, element.bbox.box_2d)
        
        base64_img = ImageUtils.encode_image(target_image_path)
        
        # 构建prompt
        system_prompt = self.get_system_prompt()
        
        context = f"""全局上下文：
- 页面核心内容：{global_summary}
- 科研假设：{hypothesis} (请特别关注数据是否支持此假设)
- 当前元素：{element.description}
- 关键词提示：{element.medical_context}

任务：
1. 提取核心数据。
2. **验证假设**：该元素展示的数据是支持、反驳还是无关该假设？"""
        
        try:
            response = self.vlm_client.chat.completions.create(
                model=self.model,
                messages=[
                    {"role": "system", "content": system_prompt},
                    {"role": "user", "content": [
                        {"type": "text", "text": context},
                        {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_img}"}}
                    ]}
                ],
                max_tokens=500
            )
            
            key_insight = response.choices[0].message.content
            
            return ElementInsight(
                element_id=element.element_id,
                element_type=element.type,
                key_insight=key_insight,
                hypothesis_verification="见核心洞察",
                data_evidence=element.description,
                confidence=0.8,
                status="pass",
                crop_path=target_image_path
            )
        except Exception as e:
            print(f"[元素Agent] [ERROR] 分析失败: {e}")
            return ElementInsight(
                element_id=element.element_id,
                element_type=element.type,
                key_insight="分析失败",
                data_evidence=element.description,
                confidence=0.0,
                status="fail",
                crop_path=target_image_path
            )


class ChartAnalysisAgent(ElementAnalysisAgent):
    """图表分析Agent"""
    def __init__(self, vlm_client: VLMClient):
        super().__init__(vlm_client, "chart")


class TableAnalysisAgent(ElementAnalysisAgent):
    """表格分析Agent"""
    def __init__(self, vlm_client: VLMClient):
        super().__init__(vlm_client, "table")
        self.textin_client = _get_textin_client()

    def _analyze_markdown_table_text_only(self,
                                          markdown_table: str,
                                          global_summary: str,
                                          hypothesis: str = "",
                                          caption: str = "",
                                          footnotes: Optional[list] = None) -> str:
        footnotes = footnotes or []
        note_text = "\n".join([f"- {x}" for x in footnotes if str(x).strip()])
        context = f"""全局上下文：
- 页面核心内容：{global_summary}
- 科研假设：{hypothesis}
- 表格标题/说明：{caption}

表格Markdown：
{markdown_table}

补充脚注：
{note_text if note_text else '(无)'}
"""

        system_prompt = """你是临床与商业数据分析专家。你将基于“结构化表格文本（非图片）”做洞察。

要求：
1. 严禁编造数据；结论必须来自给定表格。
2. 洞察必须明确引用字段名或行列位置（例如“第2行‘OS(月)’列”）。
3. 优先输出关键指标、组间差异、统计显著性（若有 p 值/HR/CI）。
4. 结果使用中文，结构化输出：
   - 关键指标
   - 主要差异
   - 与科研假设关系（支持/反驳/无关）
"""

        response = self.vlm_client.chat.completions.create(
            model=self.model,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": context}
            ],
            max_tokens=700
        )
        return response.choices[0].message.content

    def analyze_element(self,
                       element: PageElement,
                       global_summary: str,
                       original_image_path: str,
                       hypothesis: str = "",
                       native_page_evidence: Optional[dict] = None) -> ElementInsight:
        """表格元素优先使用 PPTX 原生表格，其次 TextIn，再回退 VLM。"""
        print(f"[元素Agent] 分析 table 元素(原生优先): {element.element_id}")

        # 裁剪 ROI
        target_image_path = original_image_path
        if element.bbox and element.bbox.box_2d:
            target_image_path = ImageUtils.crop_image(original_image_path, element.bbox.box_2d)

        native_tables = (native_page_evidence or {}).get("native_tables", []) if native_page_evidence else []
        native_match = _pick_native_table_for_element(element, native_tables)
        if native_match:
            md_table = (native_match.get("markdown_table") or "").strip()
            caption = (native_match.get("caption") or native_match.get("shape_name") or "").strip()
            footnotes = native_match.get("footnotes", [])
            if md_table:
                try:
                    key_insight = self._analyze_markdown_table_text_only(
                        markdown_table=md_table,
                        global_summary=global_summary,
                        hypothesis=hypothesis,
                        caption=caption,
                        footnotes=footnotes,
                    )
                    cross_check_note = ""
                    if self.textin_client:
                        try:
                            ocr_result = self.textin_client.recognize_table(target_image_path)
                            ocr_md = (ocr_result.get("markdown_table") or "").strip()
                            if ocr_md:
                                cross_check_note = f" | cross-check: TextIn可用({len(ocr_md)} chars)"
                        except Exception:
                            pass

                    return ElementInsight(
                        element_id=element.element_id,
                        element_type=element.type,
                        key_insight=key_insight,
                        hypothesis_verification=f"native_table_first{cross_check_note}",
                        data_evidence=md_table,
                        confidence=0.96,
                        status="pass",
                        crop_path=target_image_path,
                    )
                except Exception as e:
                    print(f"[元素Agent] [WARN] 原生表格洞察失败，回退OCR/VLM: {e}")

        if not self.textin_client:
            # 没有 TextIn 客户端时，降级回原逻辑（VLM 看图）
            return super().analyze_element(element, global_summary, original_image_path, hypothesis, native_page_evidence)

        try:
            result = self.textin_client.recognize_table(target_image_path)
            md_table = (result.get("markdown_table") or "").strip()

            # 如果拿到的是占位 mock 表格，则视为识别失败，回退 VLM（避免误导后续 Supervisor/输出）
            normalized = "\n".join([ln.strip() for ln in md_table.splitlines() if ln.strip()])
            is_mock = (
                "| Column1 | Column2 |" in normalized
                and "| Data1 | Data2 |" in normalized
                and len(normalized.splitlines()) <= 4
            )
            if is_mock:
                md_table = ""

            if not md_table:
                # TextIn 没拿到表格，回退 VLM
                return super().analyze_element(element, global_summary, original_image_path, hypothesis, native_page_evidence)

            # 简洁输出：给后续 Supervisor/导出足够证据
            key_insight = "\n".join([
                "识别方式：合合(TextIn)表格识别",
                "\n### 表格结构化结果（Markdown）\n",
                md_table,
            ])

            return ElementInsight(
                element_id=element.element_id,
                element_type=element.type,
                key_insight=key_insight,
                hypothesis_verification="表格已完成结构化识别（如需验证假设，可在Supervisor阶段结合上下文判断）",
                data_evidence=md_table,
                confidence=0.9,
                status="pass",
                crop_path=target_image_path,
            )
        except Exception as e:
            print(f"[元素Agent] [ERROR] TextIn 表格识别失败: {e}")
            return super().analyze_element(element, global_summary, original_image_path, hypothesis, native_page_evidence)


class DiagramAnalysisAgent(ElementAnalysisAgent):
    """流程图/结构图分析Agent"""
    def __init__(self, vlm_client: VLMClient):
        super().__init__(vlm_client, "diagram")


class ImageAnalysisAgent(ElementAnalysisAgent):
    """图片分析Agent"""
    def __init__(self, vlm_client: VLMClient):
        super().__init__(vlm_client, "image")


class FormulaAnalysisAgent(ElementAnalysisAgent):
    """公式分析Agent"""
    def __init__(self, vlm_client: VLMClient):
        super().__init__(vlm_client, "formula")


# Agent工厂
AGENT_FACTORY: Dict[str, type] = {
    "chart": ChartAnalysisAgent,
    "table": TableAnalysisAgent,
    "formula": FormulaAnalysisAgent,
    "image": ImageAnalysisAgent,
    "unknown": ElementAnalysisAgent
}


def get_agent(element_type: str, vlm_client: VLMClient) -> ElementAnalysisAgent:
    """获取对应类型的Agent"""
    element_type_norm = normalize_element_type(element_type)
    agent_class = AGENT_FACTORY.get(element_type_norm)
    
    # 如果找到对应的Agent类，直接实例化（它们都只需要vlm_client）
    if agent_class:
        return agent_class(vlm_client)
    
    # 如果没有找到对应的Agent，使用基类ElementAnalysisAgent，并传递element_type
    return ElementAnalysisAgent(vlm_client, element_type_norm)


def node_element_worker(state: dict, vlm_client: VLMClient, debug_logger=None) -> dict:
    """节点：并行元素分析worker（会被多次调用）
    
    接收的state来自Send的payload，包含：
    - element: PageElement
    - global_summary: str
    - image_path: str
    """
    if state.get("skip_worker"):
        return {"element_insights": []}

    element = state["element"]
    global_summary = state["global_summary"]
    image_path = state["image_path"]
    page_index = state.get("page_index", 0)
    hypothesis = state.get("hypothesis", "")
    native_page_evidence = state.get("native_page_evidence", {})
    
    # 记录步骤开始
    if debug_logger:
        debug_logger.log_step_start(
            page_index,
            f"Step 2: Element Worker - {element.element_id}",
            input_data={
                "page_index": page_index,
                "element_id": element.element_id,
                "element_type": element.type,
                "element_description": element.description[:100],
                "image_path": image_path
            },
            previous_step="Step 2.5: Locator"
        )
    
    # 获取对应类型的Agent
    agent = get_agent(element.type, vlm_client)
    
    # 分析元素
    insight = agent.analyze_element(
        element,
        global_summary,
        image_path,
        hypothesis,
        native_page_evidence=native_page_evidence,
    )
    
    # 记录步骤结束
    if debug_logger:
        debug_logger.log_step_end(
            page_index,
            f"Step 2: Element Worker - {element.element_id}",
            output_data={
                "element_id": insight.element_id,
                "element_type": insight.element_type,
                "confidence": insight.confidence,
                "status": insight.status,
                "key_insight_preview": insight.key_insight[:100]
            },
            status="success"
        )
    
    return {"element_insights": [insight]}
