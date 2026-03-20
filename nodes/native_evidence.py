"""PPTX 原生页面证据抽取。

抽取内容：
- native_tables: 原生可编辑表格（cell matrix + markdown + bbox）
- native_text_blocks: 文本框/说明文字（块级）
- notes: 备注页文本
"""

from __future__ import annotations

from typing import Any, Dict, List
from pathlib import Path

from pptx import Presentation


def _shape_to_norm_bbox(shape: Any, slide_width: int, slide_height: int) -> List[int] | None:
    try:
        left = int(getattr(shape, "left", 0))
        top = int(getattr(shape, "top", 0))
        width = int(getattr(shape, "width", 0))
        height = int(getattr(shape, "height", 0))
        if slide_width <= 0 or slide_height <= 0 or width <= 0 or height <= 0:
            return None

        xmin = max(0, min(1000, round(left / slide_width * 1000)))
        ymin = max(0, min(1000, round(top / slide_height * 1000)))
        xmax = max(0, min(1000, round((left + width) / slide_width * 1000)))
        ymax = max(0, min(1000, round((top + height) / slide_height * 1000)))
        if xmax <= xmin or ymax <= ymin:
            return None
        return [ymin, xmin, ymax, xmax]
    except Exception:
        return None


def _escape_md_cell(text: str) -> str:
    return (text or "").replace("|", "\\|").replace("\n", " ").replace("\r", " ").strip()


def _table_to_matrix(table: Any) -> List[List[str]]:
    matrix: List[List[str]] = []
    for row in table.rows:
        row_cells: List[str] = []
        for cell in row.cells:
            row_cells.append((cell.text or "").strip())
        matrix.append(row_cells)
    return matrix


def _matrix_to_markdown_table(matrix: List[List[str]]) -> str:
    if not matrix:
        return ""

    col_count = max((len(r) for r in matrix), default=0)
    if col_count == 0:
        return ""

    normalized = [r + [""] * (col_count - len(r)) for r in matrix]
    header = normalized[0]
    body = normalized[1:] if len(normalized) > 1 else []

    lines: List[str] = []
    lines.append("| " + " | ".join(_escape_md_cell(c) for c in header) + " |")
    lines.append("| " + " | ".join(["---"] * col_count) + " |")
    for row in body:
        lines.append("| " + " | ".join(_escape_md_cell(c) for c in row) + " |")

    return "\n".join(lines)


def extract_native_page_evidence(ppt_path: str, page_index: int) -> Dict[str, Any]:
    evidence: Dict[str, Any] = {
        "source": "pptx-native",
        "native_tables": [],
        "native_text_blocks": [],
        "notes": "",
    }

    suffix = Path(ppt_path).suffix.lower()
    if suffix not in {".pptx", ".ppt"}:
        evidence["source"] = "non-pptx"
        return evidence

    try:
        prs = Presentation(ppt_path)
        if page_index < 0 or page_index >= len(prs.slides):
            evidence["source"] = "out-of-range"
            return evidence

        slide = prs.slides[page_index]
        slide_width = int(prs.slide_width)
        slide_height = int(prs.slide_height)

        native_tables: List[Dict[str, Any]] = []
        native_text_blocks: List[Dict[str, Any]] = []

        for shape in slide.shapes:
            bbox = _shape_to_norm_bbox(shape, slide_width, slide_height)
            shape_id = str(getattr(shape, "shape_id", ""))
            shape_name = str(getattr(shape, "name", ""))

            has_table = bool(getattr(shape, "has_table", False))
            if has_table:
                try:
                    table = shape.table
                    matrix = _table_to_matrix(table)
                    md_table = _matrix_to_markdown_table(matrix)
                    if not any(any((c or "").strip() for c in row) for row in matrix):
                        continue
                    native_tables.append(
                        {
                            "table_id": f"native_table_{shape_id or len(native_tables) + 1}",
                            "shape_id": shape_id,
                            "shape_name": shape_name,
                            "bbox": bbox,
                            "row_count": len(matrix),
                            "col_count": max((len(r) for r in matrix), default=0),
                            "cell_matrix": matrix,
                            "markdown_table": md_table,
                            "caption": "",
                            "footnotes": [],
                        }
                    )
                except Exception:
                    continue
                continue

            has_text_frame = bool(getattr(shape, "has_text_frame", False))
            if has_text_frame:
                text = ""
                try:
                    text = (shape.text_frame.text or "").strip()
                except Exception:
                    text = ""
                if text:
                    native_text_blocks.append(
                        {
                            "block_id": f"native_text_{shape_id or len(native_text_blocks) + 1}",
                            "shape_id": shape_id,
                            "shape_name": shape_name,
                            "bbox": bbox,
                            "text": text,
                        }
                    )

        notes = ""
        try:
            if slide.has_notes_slide and slide.notes_slide and slide.notes_slide.notes_text_frame:
                notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes = ""

        evidence["native_tables"] = native_tables
        evidence["native_text_blocks"] = native_text_blocks
        evidence["notes"] = notes
        return evidence

    except Exception as e:
        evidence["source"] = "error"
        evidence["error"] = str(e)
        return evidence
